"""Enhanced document viewer with multiple format support (binary, hex, image, pdf, ppt, xlsx, csv, json)."""

from __future__ import annotations

import base64
import csv
import io
import json
from pathlib import Path
from typing import Any, Dict, Optional

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response, StreamingResponse
from pydantic import BaseModel

from backend_api.routers.auth import get_current_user
from backend_api.routers.documents import _load_metadata, _file_path_from_metadata

router = APIRouter()


class DocumentViewResponse(BaseModel):
    document_id: str
    view_type: str
    content: Any
    metadata: Dict[str, Any]


def get_binary_view(file_path: Path) -> Dict[str, Any]:
    """Get binary representation of file."""
    data = file_path.read_bytes()
    return {
        "size": len(data),
        "preview": base64.b64encode(data[:1024]).decode("ascii"),  # First 1KB
        "total_size": len(data)
    }


def get_hex_view(file_path: Path, offset: int = 0, length: int = 1024) -> Dict[str, Any]:
    """Get hexadecimal representation of file."""
    data = file_path.read_bytes()
    total_size = len(data)
    
    chunk = data[offset:offset + length]
    hex_lines = []
    ascii_lines = []
    
    for i in range(0, len(chunk), 16):
        hex_part = " ".join(f"{b:02x}" for b in chunk[i:i + 16])
        ascii_part = "".join(chr(b) if 32 <= b < 127 else "." for b in chunk[i:i + 16])
        hex_lines.append(f"{offset + i:08x}: {hex_part:<48} {ascii_part}")
        ascii_lines.append(ascii_part)
    
    return {
        "offset": offset,
        "length": length,
        "total_size": total_size,
        "hex_dump": hex_lines,
        "has_more": offset + length < total_size
    }


def get_image_view(file_path: Path) -> Dict[str, Any]:
    """Get image file information and base64 encoded content."""
    try:
        from PIL import Image
        
        img = Image.open(file_path)
        img_format = img.format or "UNKNOWN"
        width, height = img.size
        mode = img.mode
        
        # Convert to base64
        img_bytes = io.BytesIO()
        img.save(img_bytes, format=img_format)
        img_base64 = base64.b64encode(img_bytes.getvalue()).decode("ascii")
        
        return {
            "format": img_format,
            "width": width,
            "height": height,
            "mode": mode,
            "base64_data": f"data:image/{img_format.lower()};base64,{img_base64}"
        }
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Not a valid image file: {str(e)}")


def get_pdf_view(file_path: Path) -> Dict[str, Any]:
    """Get PDF file information and extract text."""
    try:
        import PyPDF2
        
        with open(file_path, "rb") as f:
            pdf_reader = PyPDF2.PdfReader(f)
            num_pages = len(pdf_reader.pages)
            
            # Extract text from first few pages
            text_content = []
            for i, page in enumerate(pdf_reader.pages[:5]):  # First 5 pages
                text_content.append({
                    "page": i + 1,
                    "text": page.extract_text()[:5000]  # Limit per page
                })
            
            # Get metadata
            metadata = pdf_reader.metadata or {}
            
            return {
                "num_pages": num_pages,
                "text_content": text_content,
                "metadata": {
                    "title": metadata.get("/Title", ""),
                    "author": metadata.get("/Author", ""),
                    "subject": metadata.get("/Subject", ""),
                    "creator": metadata.get("/Creator", "")
                }
            }
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to read PDF: {str(e)}")


def get_powerpoint_view(file_path: Path) -> Dict[str, Any]:
    """Get PowerPoint file information."""
    try:
        from pptx import Presentation
        
        prs = Presentation(file_path)
        num_slides = len(prs.slides)
        
        slides_content = []
        for i, slide in enumerate(prs.slides[:10]):  # First 10 slides
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    slide_text.append(shape.text)
            
            slides_content.append({
                "slide_number": i + 1,
                "text": "\n".join(slide_text)
            })
        
        return {
            "num_slides": num_slides,
            "slides_content": slides_content
        }
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to read PowerPoint: {str(e)}")


def get_excel_view(file_path: Path, sheet_index: int = 0) -> Dict[str, Any]:
    """Get Excel file information and content."""
    try:
        import pandas as pd
        
        # Read Excel file
        excel_file = pd.ExcelFile(file_path)
        sheet_names = excel_file.sheet_names
        
        if sheet_index >= len(sheet_names):
            sheet_index = 0
        
        sheet_name = sheet_names[sheet_index]
        df = pd.read_excel(file_path, sheet_name=sheet_name, nrows=100)  # First 100 rows
        
        return {
            "sheet_names": sheet_names,
            "current_sheet": sheet_name,
            "num_rows": len(df),
            "num_columns": len(df.columns),
            "columns": list(df.columns),
            "data": df.head(50).to_dict(orient="records"),  # First 50 rows
            "has_more": len(df) > 50
        }
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to read Excel: {str(e)}")


def get_csv_view(file_path: Path, delimiter: Optional[str] = None) -> Dict[str, Any]:
    """Get CSV file information and content."""
    try:
        import pandas as pd
        
        # Try to detect delimiter
        if delimiter is None:
            sample = file_path.read_text(encoding="utf-8", errors="ignore")[:1024]
            sniffer = csv.Sniffer()
            delimiter = sniffer.sniff(sample).delimiter
        
        df = pd.read_csv(file_path, delimiter=delimiter, nrows=100)
        
        return {
            "delimiter": delimiter,
            "num_rows": len(df),
            "num_columns": len(df.columns),
            "columns": list(df.columns),
            "data": df.head(50).to_dict(orient="records"),
            "has_more": len(df) > 50
        }
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to read CSV: {str(e)}")


def get_json_view(file_path: Path) -> Dict[str, Any]:
    """Get JSON file content with formatting."""
    try:
        content = file_path.read_text(encoding="utf-8")
        data = json.loads(content)
        
        return {
            "formatted": json.dumps(data, indent=2),
            "minified": json.dumps(data, separators=(",", ":")),
            "type": type(data).__name__,
            "size": len(content)
        }
    except json.JSONDecodeError as e:
        raise HTTPException(status_code=400, detail=f"Invalid JSON: {str(e)}")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to read JSON: {str(e)}")


def get_text_view(file_path: Path) -> Dict[str, Any]:
    """Get text file content."""
    try:
        content = file_path.read_text(encoding="utf-8", errors="ignore")
        return {
            "content": content,
            "size": len(content),
            "lines": content.count("\n") + 1
        }
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to read text file: {str(e)}")


@router.get("/documents/{document_id}/view/{view_type}")
async def get_document_view(
    document_id: str,
    view_type: str,
    offset: Optional[int] = None,
    length: Optional[int] = None,
    sheet_index: Optional[int] = None,
    delimiter: Optional[str] = None,
    current_user: dict = Depends(get_current_user)
):
    """Get document in specified view format."""
    doc = _load_metadata(document_id)
    file_path = _file_path_from_metadata(doc)
    
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="Document file not found")
    
    ext = file_path.suffix.lower()
    
    # Route to appropriate viewer
    if view_type == "binary":
        content = get_binary_view(file_path)
    elif view_type == "hex":
        content = get_hex_view(file_path, offset=offset or 0, length=length or 1024)
    elif view_type == "image":
        if ext not in [".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".svg"]:
            raise HTTPException(status_code=400, detail="File is not an image")
        content = get_image_view(file_path)
    elif view_type == "pdf":
        if ext != ".pdf":
            raise HTTPException(status_code=400, detail="File is not a PDF")
        content = get_pdf_view(file_path)
    elif view_type == "powerpoint" or view_type == "ppt":
        if ext not in [".ppt", ".pptx"]:
            raise HTTPException(status_code=400, detail="File is not a PowerPoint file")
        content = get_powerpoint_view(file_path)
    elif view_type == "excel" or view_type == "xlsx":
        if ext not in [".xls", ".xlsx", ".xlsm", ".xlsb"]:
            raise HTTPException(status_code=400, detail="File is not an Excel file")
        content = get_excel_view(file_path, sheet_index=sheet_index or 0)
    elif view_type == "csv":
        if ext != ".csv":
            raise HTTPException(status_code=400, detail="File is not a CSV file")
        content = get_csv_view(file_path, delimiter=delimiter)
    elif view_type == "json":
        if ext != ".json":
            raise HTTPException(status_code=400, detail="File is not a JSON file")
        content = get_json_view(file_path)
    elif view_type == "text":
        content = get_text_view(file_path)
    else:
        raise HTTPException(status_code=400, detail=f"Unsupported view type: {view_type}")
    
    return DocumentViewResponse(
        document_id=document_id,
        view_type=view_type,
        content=content,
        metadata={
            "filename": doc["filename"],
            "file_type": doc.get("file_type", ""),
            "category": doc.get("category", ""),
            "size_bytes": file_path.stat().st_size
        }
    )


@router.get("/documents/{document_id}/view/{view_type}/download")
async def download_document_view(
    document_id: str,
    view_type: str,
    current_user: dict = Depends(get_current_user)
):
    """Download document in specified format."""
    doc = _load_metadata(document_id)
    file_path = _file_path_from_metadata(doc)
    
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="Document file not found")
    
    if view_type == "hex":
        content = get_hex_view(file_path, offset=0, length=file_path.stat().st_size)
        hex_content = "\n".join(content["hex_dump"])
        return Response(
            content=hex_content,
            media_type="text/plain",
            headers={"Content-Disposition": f'attachment; filename="{doc["filename"]}.hex"'}
        )
    else:
        # Return original file
        return StreamingResponse(
            open(file_path, "rb"),
            media_type="application/octet-stream",
            headers={"Content-Disposition": f'attachment; filename="{doc["filename"]}"'}
        )
