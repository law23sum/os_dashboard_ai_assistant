"""
File processors for converting common office formats to and from CIR.

All dependencies are optional; when a library is missing the processor will
raise a clear ImportError so callers can decide to fall back to a simpler path.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional

from api_connectors.universal_connector import (
    CIRDocument,
    ContentBlock,
    ContentBlockType,
    Section,
)

# Optional imports guarded so this module remains importable without extras.
try:  # pragma: no cover - optional dependency
    from docx import Document  # type: ignore
except Exception:  # pragma: no cover - optional dependency
    Document = None

try:  # pragma: no cover - optional dependency
    import openpyxl  # type: ignore
except Exception:  # pragma: no cover - optional dependency
    openpyxl = None

try:  # pragma: no cover - optional dependency
    from pptx import Presentation  # type: ignore
except Exception:  # pragma: no cover - optional dependency
    Presentation = None

try:  # pragma: no cover - optional dependency
    import fitz  # type: ignore
    from PIL import Image  # type: ignore
except Exception:  # pragma: no cover - optional dependency
    fitz = None
    Image = None


class MissingDependencyError(ImportError):
    """Raised when a processor dependency is not installed."""


def _require(dep, message: str):
    if dep is None:
        raise MissingDependencyError(message)


class WordParser:
    async def process_file(self, path: Path) -> CIRDocument:
        _require(Document, "python-docx is required for Word parsing")
        doc = Document(str(path))
        cir = CIRDocument(title=doc.core_properties.title or path.stem, document_type="word")
        for i, paragraph in enumerate(doc.paragraphs):
            cir.sections.append(
                Section(
                    title=f"Paragraph {i+1}",
                    content_blocks=[ContentBlock(ContentBlockType.TEXT, paragraph.text)],
                )
            )
        return cir

    async def generate_file(self, cir: CIRDocument, path: Path) -> None:
        _require(Document, "python-docx is required for Word generation")
        doc = Document()
        if cir.title:
            doc.add_heading(cir.title, level=1)
        for section in cir.sections:
            for block in section.content_blocks:
                if block.block_type == ContentBlockType.TEXT:
                    doc.add_paragraph(str(block.content))
                elif block.block_type == ContentBlockType.TABLE:
                    table = doc.add_table(rows=len(block.content.get("rows", [])), cols=len(block.content.get("headers", [])) or 1)
                    rows = block.content.get("rows", [])
                    for r_idx, row in enumerate(rows):
                        for c_idx, cell in enumerate(row):
                            table.cell(r_idx, c_idx).text = str(cell)
        path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(path))


class ExcelParser:
    async def process_file(self, path: Path) -> CIRDocument:
        _require(openpyxl, "openpyxl is required for Excel parsing")
        wb = openpyxl.load_workbook(str(path), data_only=True)
        cir = CIRDocument(title=wb.properties.title or path.stem, document_type="excel")
        for ws in wb.worksheets:
            rows: List[List[str]] = []
            for row in ws.iter_rows(values_only=True):
                rows.append([("" if cell is None else cell) for cell in row])
            cir.sections.append(
                Section(
                    title=ws.title,
                    content_blocks=[ContentBlock(ContentBlockType.TABLE, {"rows": rows})],
                )
            )
        return cir

    async def generate_file(self, cir: CIRDocument, path: Path) -> None:
        _require(openpyxl, "openpyxl is required for Excel generation")
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for section in cir.sections:
            ws = wb.create_sheet(title=section.title[:31] or "Sheet")
            for block in section.content_blocks:
                if block.block_type == ContentBlockType.TABLE:
                    for r_idx, row in enumerate(block.content.get("rows", []), start=1):
                        for c_idx, cell in enumerate(row, start=1):
                            ws.cell(row=r_idx, column=c_idx, value=cell)
                elif block.block_type == ContentBlockType.TEXT:
                    ws.cell(row=1, column=1, value=str(block.content))
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(path))


class PowerPointParser:
    async def process_file(self, path: Path) -> CIRDocument:
        _require(Presentation, "python-pptx is required for PowerPoint parsing")
        prs = Presentation(str(path))
        cir = CIRDocument(title=prs.core_properties.title or path.stem, document_type="powerpoint")
        for idx, slide in enumerate(prs.slides, start=1):
            blocks: List[ContentBlock] = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    blocks.append(ContentBlock(ContentBlockType.TEXT, shape.text))
            cir.sections.append(Section(title=f"Slide {idx}", content_blocks=blocks))
        return cir

    async def generate_file(self, cir: CIRDocument, path: Path) -> None:
        _require(Presentation, "python-pptx is required for PowerPoint generation")
        prs = Presentation()
        blank = prs.slide_layouts[6]
        for section in cir.sections:
            slide = prs.slides.add_slide(blank)
            body = slide.shapes.add_textbox(prs.slide_width * 0.05, prs.slide_height * 0.05, prs.slide_width * 0.9, prs.slide_height * 0.9)
            tf = body.text_frame
            tf.clear()
            for block in section.content_blocks:
                if block.block_type == ContentBlockType.TEXT:
                    p = tf.add_paragraph()
                    p.text = str(block.content)
        path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(path))


class PdfParser:
    def __init__(self, ocr=None):
        self.ocr = ocr

    async def process_file(self, path: Path) -> CIRDocument:
        _require(fitz, "PyMuPDF (fitz) is required for PDF parsing")
        doc = fitz.open(str(path))
        cir = CIRDocument(title=doc.metadata.get("title") or path.stem, document_type="pdf")
        for i, page in enumerate(doc, start=1):
            text = page.get_text()
            if not text.strip() and self.ocr and Image is not None:
                pix = page.get_pixmap()
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                text = self.ocr.image_to_string(img)
            cir.sections.append(Section(title=f"Page {i}", content_blocks=[ContentBlock(ContentBlockType.TEXT, text)]))
        doc.close()
        return cir

    async def generate_file(self, cir: CIRDocument, path: Path) -> None:
        _require(fitz, "PyMuPDF (fitz) is required for PDF generation")
        doc = fitz.open()
        for section in cir.sections:
            page = doc.new_page()
            y = 50
            page.insert_text((50, y), section.title, fontsize=16)
            y += 24
            for block in section.content_blocks:
                if block.block_type == ContentBlockType.TEXT:
                    page.insert_text((50, y), str(block.content), fontsize=11)
                    y += 14 * (str(block.content).count("\n") + 1)
        path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(path))
        doc.close()


@dataclass
class FileProcessorFactory:
    """Factory to build processor mapping based on installed deps."""

    base_path: Path

    def build(self) -> Dict[str, object]:
        processors: Dict[str, object] = {}
        if Document:
            processors[".docx"] = WordParser()
        if openpyxl:
            processors[".xlsx"] = ExcelParser()
        if Presentation:
            processors[".pptx"] = PowerPointParser()
        if fitz:
            processors[".pdf"] = PdfParser()
        return processors

