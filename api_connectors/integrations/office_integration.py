"""
Office Document Integration

Handles Word, Excel, and PowerPoint file processing
"""

import os
from typing import Dict, Any, Optional, List, Tuple
from pathlib import Path
import logging

try:
    from docx import Document
    from docx.shared import Inches
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    import openpyxl
    from openpyxl import Workbook
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

try:
    from pptx import Presentation
    from pptx.util import Inches as PptxInches
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

from config.logging_config import setup_logger

logger = setup_logger(__name__)


class OfficeIntegration:
    """Office document processing integration"""

    def __init__(self):
        self.logger = setup_logger("OfficeIntegration")

    def is_available(self) -> Dict[str, bool]:
        """Check which Office libraries are available"""
        return {
            "word": DOCX_AVAILABLE,
            "excel": OPENPYXL_AVAILABLE,
            "powerpoint": PPTX_AVAILABLE
        }

    # Word Document Operations

    def read_word_document(self, file_path: str) -> Dict[str, Any]:
        """Read Word document content"""
        try:
            if not DOCX_AVAILABLE:
                return {"error": "python-docx not available"}

            doc = Document(file_path)

            content = {
                "paragraphs": [],
                "tables": [],
                "images": [],
                "metadata": {}
            }

            # Extract paragraphs
            for para in doc.paragraphs:
                if para.text.strip():
                    content["paragraphs"].append({
                        "text": para.text,
                        "style": para.style.name if para.style else None
                    })

            # Extract tables
            for table in doc.tables:
                table_data = []
                for row in table.rows:
                    row_data = [cell.text for cell in row.cells]
                    table_data.append(row_data)
                content["tables"].append(table_data)

            # Extract metadata
            core_props = doc.core_properties
            content["metadata"] = {
                "title": core_props.title,
                "author": core_props.author,
                "created": core_props.created.isoformat() if core_props.created else None,
                "modified": core_props.modified.isoformat() if core_props.modified else None
            }

            return content

        except Exception as e:
            self.logger.error(f"Word document reading failed: {e}")
            return {"error": str(e)}

    def create_word_document(self, content: Dict[str, Any], output_path: str) -> bool:
        """Create Word document"""
        try:
            if not DOCX_AVAILABLE:
                return False

            doc = Document()

            # Add title
            if "title" in content:
                doc.add_heading(content["title"], 0)

            # Add paragraphs
            if "paragraphs" in content:
                for para in content["paragraphs"]:
                    if isinstance(para, str):
                        doc.add_paragraph(para)
                    elif isinstance(para, dict):
                        p = doc.add_paragraph(para.get("text", ""))
                        if "style" in para:
                            p.style = para["style"]

            # Add tables
            if "tables" in content:
                for table_data in content["tables"]:
                    if table_data:
                        table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
                        for i, row_data in enumerate(table_data):
                            for j, cell_data in enumerate(row_data):
                                table.cell(i, j).text = str(cell_data)

            doc.save(output_path)
            return True

        except Exception as e:
            self.logger.error(f"Word document creation failed: {e}")
            return False

    def extract_text_from_word(self, file_path: str) -> str:
        """Extract plain text from Word document"""
        try:
            if not DOCX_AVAILABLE:
                return ""

            doc = Document(file_path)
            text = []

            for para in doc.paragraphs:
                if para.text.strip():
                    text.append(para.text)

            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text.append(cell.text)

            return "\n".join(text)

        except Exception as e:
            self.logger.error(f"Word text extraction failed: {e}")
            return ""

    # Excel Operations

    def read_excel_file(self, file_path: str, sheet_name: Optional[str] = None) -> Dict[str, Any]:
        """Read Excel file content"""
        try:
            if not OPENPYXL_AVAILABLE:
                return {"error": "openpyxl not available"}

            wb = openpyxl.load_workbook(file_path, data_only=True)

            content = {
                "sheets": {},
                "metadata": {
                    "sheet_names": wb.sheetnames,
                    "active_sheet": wb.active.title
                }
            }

            # Read all sheets or specific sheet
            sheets_to_read = [sheet_name] if sheet_name else wb.sheetnames

            for sheet_name in sheets_to_read:
                if sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]

                    # Get data range
                    data = []
                    for row in ws.iter_rows(values_only=True):
                        data.append(list(row))

                    # Remove empty rows at end
                    while data and all(cell is None for cell in data[-1]):
                        data.pop()

                    content["sheets"][sheet_name] = {
                        "data": data,
                        "dimensions": f"{ws.max_row}x{ws.max_column}",
                        "merged_cells": [str(range) for range in ws.merged_cells.ranges]
                    }

            return content

        except Exception as e:
            self.logger.error(f"Excel reading failed: {e}")
            return {"error": str(e)}

    def create_excel_file(self, data: Dict[str, Any], output_path: str) -> bool:
        """Create Excel file"""
        try:
            if not OPENPYXL_AVAILABLE:
                return False

            wb = Workbook()

            # Remove default sheet
            wb.remove(wb.active)

            for sheet_name, sheet_data in data.get("sheets", {}).items():
                ws = wb.create_sheet(sheet_name)

                # Add data
                for row_idx, row_data in enumerate(sheet_data.get("data", []), 1):
                    for col_idx, cell_value in enumerate(row_data, 1):
                        ws.cell(row=row_idx, column=col_idx, value=cell_value)

            wb.save(output_path)
            return True

        except Exception as e:
            self.logger.error(f"Excel creation failed: {e}")
            return False

    def get_excel_summary(self, file_path: str) -> Dict[str, Any]:
        """Get Excel file summary"""
        try:
            if not OPENPYXL_AVAILABLE:
                return {"error": "openpyxl not available"}

            wb = openpyxl.load_workbook(file_path, read_only=True)

            summary = {
                "sheet_count": len(wb.sheetnames),
                "sheet_names": wb.sheetnames,
                "total_cells": 0,
                "non_empty_cells": 0
            }

            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                cell_count = 0
                non_empty_count = 0

                for row in ws.iter_rows():
                    for cell in row:
                        cell_count += 1
                        if cell.value is not None:
                            non_empty_count += 1

                summary["total_cells"] += cell_count
                summary["non_empty_cells"] += non_empty_count

            return summary

        except Exception as e:
            self.logger.error(f"Excel summary failed: {e}")
            return {"error": str(e)}

    # PowerPoint Operations

    def read_powerpoint_file(self, file_path: str) -> Dict[str, Any]:
        """Read PowerPoint file content"""
        try:
            if not PPTX_AVAILABLE:
                return {"error": "python-pptx not available"}

            prs = Presentation(file_path)

            content = {
                "slides": [],
                "metadata": {
                    "slide_count": len(prs.slides),
                    "slide_size": f"{prs.slide_width/914400:.1f}x{prs.slide_height/914400:.1f} inches"
                }
            }

            for slide_idx, slide in enumerate(prs.slides):
                slide_content = {
                    "slide_number": slide_idx + 1,
                    "title": None,
                    "content": [],
                    "notes": ""
                }

                # Extract shapes
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        if not slide_content["title"]:
                            slide_content["title"] = shape.text
                        else:
                            slide_content["content"].append(shape.text)

                # Extract notes
                if slide.notes_slide and slide.notes_slide.notes_text_frame:
                    slide_content["notes"] = slide.notes_slide.notes_text_frame.text

                content["slides"].append(slide_content)

            return content

        except Exception as e:
            self.logger.error(f"PowerPoint reading failed: {e}")
            return {"error": str(e)}

    def create_powerpoint_file(self, content: Dict[str, Any], output_path: str) -> bool:
        """Create PowerPoint file"""
        try:
            if not PPTX_AVAILABLE:
                return False

            prs = Presentation()

            for slide_data in content.get("slides", []):
                slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and content layout

                # Add title
                if slide_data.get("title"):
                    title = slide.shapes.title
                    title.text = slide_data["title"]

                # Add content
                if slide_data.get("content"):
                    content_shape = slide.shapes.placeholders[1]
                    content_shape.text = "\n".join(slide_data["content"])

            prs.save(output_path)
            return True

        except Exception as e:
            self.logger.error(f"PowerPoint creation failed: {e}")
            return False

    def extract_text_from_powerpoint(self, file_path: str) -> str:
        """Extract text from PowerPoint file"""
        try:
            if not PPTX_AVAILABLE:
                return ""

            prs = Presentation(file_path)
            text_parts = []

            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        text_parts.append(shape.text)

                # Extract notes
                if slide.notes_slide and slide.notes_slide.notes_text_frame:
                    text_parts.append(slide.notes_slide.notes_text_frame.text)

            return "\n".join(text_parts)

        except Exception as e:
            self.logger.error(f"PowerPoint text extraction failed: {e}")
            return ""

    # Utility Methods

    def get_file_info(self, file_path: str) -> Dict[str, Any]:
        """Get file information"""
        try:
            path = Path(file_path)
            stat = path.stat()

            return {
                "filename": path.name,
                "extension": path.suffix.lower(),
                "size": stat.st_size,
                "modified": stat.st_mtime,
                "type": self._detect_file_type(path.suffix.lower())
            }

        except Exception as e:
            return {"error": str(e)}

    def _detect_file_type(self, extension: str) -> str:
        """Detect file type from extension"""
        type_map = {
            ".docx": "word",
            ".xlsx": "excel",
            ".pptx": "powerpoint",
            ".doc": "word_legacy",
            ".xls": "excel_legacy",
            ".ppt": "powerpoint_legacy"
        }
        return type_map.get(extension, "unknown")

    def convert_office_to_pdf(self, input_path: str, output_path: str) -> bool:
        """Convert Office document to PDF (requires external tools)"""
        try:
            # This would require external tools like LibreOffice or Microsoft Office
            # For now, return False as it's not implemented
            self.logger.warning("Office to PDF conversion not implemented - requires external tools")
            return False
        except Exception as e:
            self.logger.error(f"PDF conversion failed: {e}")
            return False
