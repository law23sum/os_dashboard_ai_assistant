"""Word service for drafting and rewriting documents with AI."""

from __future__ import annotations

from pathlib import Path
from typing import Dict, List

from docx import Document

from ...ai_layer.tools import summarize_text
from ...versioning import enqueue_git_commit
from .cloud_client import CloudWordClient
from .local_client import LocalDocument, load_document, save_document


class WordService:
    """High-level Word document operations."""

    def draft_summary(self, source_text: str, destination: str, actor: str = "Aria") -> List[str]:
        summary = summarize_text(source_text, style="clear")
        doc = Document()
        doc.add_paragraph(summary)
        dest_path = Path(destination)
        dest_path.parent.mkdir(parents=True, exist_ok=True)
        save_document(doc, destination)
        # Auto-commit the new document
        from ...versioning import enqueue_commit
        enqueue_commit([str(dest_path)], actor=actor, tag="word", reason="Draft summary document")
        return [str(dest_path)]

    def rewrite_document(self, path: str, style: str = "concise", actor: str = "Aria") -> List[str]:
        doc = load_document(path)
        paragraphs = "\n".join(p.text for p in doc.paragraphs)
        rewritten = summarize_text(paragraphs, style=style)
        new_doc = Document()
        for line in rewritten.split("\n"):
            new_doc.add_paragraph(line)
        save_document(new_doc, path)
        # Auto-commit the rewritten document
        from ...versioning import enqueue_commit
        enqueue_commit([path], actor=actor, tag="word", reason=f"Rewrite document ({style})")
        return [path]


# Backwards compatibility functions
def draft_local_revision(path: str, draft: str, *, actor: str = "Aria", note: str = "") -> Dict[str, str]:
    """Write a revision text file next to an existing document and auto-commit (backwards compatibility)."""
    document = LocalDocument(path)
    revision_path = document.write_revision(draft, note=note or f"Drafted by {actor}")
    enqueue_git_commit([revision_path], actor=actor, reason="Word draft", tag="word")
    return {"document": str(document.path), "revision_path": str(revision_path)}


def upload_cloud_revision(client: CloudWordClient, drive_item_id: str, content: bytes, *, actor: str = "Aria") -> Dict[str, str]:
    """Upload a new document version to OneDrive/SharePoint and record metadata (backwards compatibility)."""
    response = client.upload_document(drive_item_id, content)
    return {"drive_item_id": drive_item_id, "status": response.get("id", "uploaded"), "actor": actor}
