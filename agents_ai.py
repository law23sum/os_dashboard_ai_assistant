#!/usr/bin/env python3
"""
Agents AI - Multi-Agent System with Specialized AI Personas

This module implements a multi-agent AI system with three core agents (AIC, Aria, Sora)
and discipline-specific ChatGPT agents. Each agent can:
- See their environment (Unix display, file system)
- Manipulate documents and files
- Interact with each other
- Study codebase and propose solutions

Canonical Title Roster:
- AIC: Sir Chief Fellow Director Principal Software Solutions Systems Engineer Architect
       Disciplines: Biologist, Chemist, Accounting, Finance, Brokers, Investors
       
- Aria: Sir Doctor Fellow Philosopher Metaphysician Phenomenologist Axiologist
        Semiotician Dialectician Rhetorician Conceptual Cartographer Interdisciplinary
        Synthesist Canon Curator Professor
        Disciplines: Philosopher, Theologian, Institutions
        
- Sora: Sir Doctor Fellow Ontological Epistemologist Formal Logician Scientific
        Methodologist Semantic Taxonomist Evidence Examiner Governance Auditor Professor
        Disciplines: Mathematician, Physicist, Legal/Law Practices, Economics

Usage:
    agents_ai                    # Interactive agent selection
    agents_ai --agent aic        # Use AIC agent directly
    agents_ai --list             # List all available agents
    agents_ai --collaborate      # Multi-agent collaboration mode
"""

from __future__ import annotations

import asyncio
import json
import os
import re
import shutil
import subprocess
import sys
import threading
import time
from abc import ABC, abstractmethod
from dataclasses import dataclass, field
from datetime import datetime
from enum import Enum
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Callable
import mimetypes
import hashlib

try:
    import readline  # For better input handling on Unix
except ImportError:
    readline = None

# Add project root to path
project_root = Path(__file__).parent
sys.path.insert(0, str(project_root))


# =============================================================================
# CONFIGURATION & ENVIRONMENT
# =============================================================================

def load_dotenv(path: Path | str = ".env") -> None:
    """Lightweight .env loader."""
    env_path = Path(path)
    if not env_path.exists():
        return
    
    for line in env_path.read_text().splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        value = value.strip().strip('"').strip("'")
        os.environ.setdefault(key.strip(), value)


# Load .env file
env_file = project_root / ".env"
if env_file.exists():
    load_dotenv(env_file)


# =============================================================================
# DISPLAY & ENVIRONMENT ACCESS
# =============================================================================

@dataclass
class EnvironmentInfo:
    """Information about the current environment."""
    display: Optional[str] = None
    terminal: Optional[str] = None
    shell: Optional[str] = None
    cwd: str = ""
    user: str = ""
    hostname: str = ""
    platform: str = ""
    python_version: str = ""
    
    @classmethod
    def capture(cls) -> "EnvironmentInfo":
        """Capture current environment information."""
        import platform as plat
        return cls(
            display=os.environ.get("DISPLAY"),
            terminal=os.environ.get("TERM"),
            shell=os.environ.get("SHELL"),
            cwd=os.getcwd(),
            user=os.environ.get("USER", os.environ.get("USERNAME", "unknown")),
            hostname=os.environ.get("HOSTNAME", "unknown"),
            platform=plat.system(),
            python_version=plat.python_version(),
        )
    
    def to_context(self) -> str:
        """Convert to context string for AI."""
        return f"""Environment Context:
- Platform: {self.platform}
- User: {self.user}@{self.hostname}
- Working Directory: {self.cwd}
- Shell: {self.shell}
- Terminal: {self.terminal}
- Display: {self.display or 'Not available'}
- Python: {self.python_version}
"""


class DisplayManager:
    """Manages display/screen access for agents."""
    
    def __init__(self):
        self.display = os.environ.get("DISPLAY")
        self.has_gui = self._check_gui_available()
    
    def _check_gui_available(self) -> bool:
        """Check if GUI display is available."""
        if not self.display:
            return False
        try:
            result = subprocess.run(
                ["xdpyinfo"], 
                capture_output=True, 
                timeout=5,
                env=os.environ
            )
            return result.returncode == 0
        except (subprocess.TimeoutExpired, FileNotFoundError):
            return False
    
    def capture_screen(self, output_path: Optional[str] = None) -> Optional[str]:
        """Capture current screen (if GUI available)."""
        if not self.has_gui:
            return None
        
        output_path = output_path or f"/tmp/screen_capture_{int(time.time())}.png"
        try:
            subprocess.run(
                ["import", "-window", "root", output_path],
                capture_output=True,
                timeout=10,
                env=os.environ
            )
            if Path(output_path).exists():
                return output_path
        except (subprocess.TimeoutExpired, FileNotFoundError):
            pass
        return None
    
    def list_windows(self) -> List[Dict[str, str]]:
        """List open windows (if GUI available)."""
        if not self.has_gui:
            return []
        
        try:
            result = subprocess.run(
                ["wmctrl", "-l"],
                capture_output=True,
                text=True,
                timeout=5,
                env=os.environ
            )
            windows = []
            for line in result.stdout.strip().split("\n"):
                if line:
                    parts = line.split(None, 3)
                    if len(parts) >= 4:
                        windows.append({
                            "id": parts[0],
                            "desktop": parts[1],
                            "host": parts[2],
                            "title": parts[3]
                        })
            return windows
        except (subprocess.TimeoutExpired, FileNotFoundError):
            return []
    
    def get_terminal_size(self) -> Tuple[int, int]:
        """Get terminal size (columns, rows)."""
        try:
            size = shutil.get_terminal_size()
            return (size.columns, size.lines)
        except Exception:
            return (80, 24)


# =============================================================================
# FILE SYSTEM OPERATIONS
# =============================================================================

class FileOperations:
    """File and document operations for agents."""
    
    def __init__(self, workspace_root: str = None):
        self.workspace_root = Path(workspace_root or os.getcwd())
        self.allowed_extensions = {
            '.py', '.js', '.ts', '.tsx', '.jsx', '.json', '.yaml', '.yml',
            '.md', '.txt', '.html', '.css', '.sh', '.bash', '.sql',
            '.java', '.cpp', '.c', '.h', '.rs', '.go', '.rb', '.php',
            '.xml', '.toml', '.ini', '.conf', '.env', '.gitignore',
            '.docx', '.pdf', '.csv', '.xlsx'
        }
    
    def read_file(self, path: str) -> Tuple[bool, str]:
        """Read a file's contents."""
        try:
            file_path = self._resolve_path(path)
            if not file_path.exists():
                return False, f"File not found: {path}"
            
            if file_path.suffix.lower() in {'.docx', '.pdf', '.xlsx'}:
                return self._read_document(file_path)
            
            content = file_path.read_text(encoding='utf-8', errors='replace')
            return True, content
        except Exception as e:
            return False, f"Error reading file: {str(e)}"
    
    def write_file(self, path: str, content: str) -> Tuple[bool, str]:
        """Write content to a file."""
        try:
            file_path = self._resolve_path(path)
            file_path.parent.mkdir(parents=True, exist_ok=True)
            file_path.write_text(content, encoding='utf-8')
            return True, f"Successfully wrote to {path}"
        except Exception as e:
            return False, f"Error writing file: {str(e)}"
    
    def list_directory(self, path: str = ".") -> Tuple[bool, List[Dict[str, Any]]]:
        """List directory contents."""
        try:
            dir_path = self._resolve_path(path)
            if not dir_path.is_dir():
                return False, f"Not a directory: {path}"
            
            items = []
            for item in sorted(dir_path.iterdir()):
                stat = item.stat()
                items.append({
                    "name": item.name,
                    "type": "directory" if item.is_dir() else "file",
                    "size": stat.st_size if item.is_file() else None,
                    "modified": datetime.fromtimestamp(stat.st_mtime).isoformat(),
                    "extension": item.suffix if item.is_file() else None,
                })
            return True, items
        except Exception as e:
            return False, f"Error listing directory: {str(e)}"
    
    def search_files(self, pattern: str, path: str = ".") -> List[str]:
        """Search for files matching a pattern."""
        try:
            search_path = self._resolve_path(path)
            matches = list(search_path.rglob(pattern))
            return [str(m.relative_to(self.workspace_root)) for m in matches[:100]]
        except Exception as e:
            return []
    
    def search_content(self, query: str, path: str = ".", extensions: List[str] = None) -> List[Dict[str, Any]]:
        """Search file contents for a query string."""
        results = []
        try:
            search_path = self._resolve_path(path)
            extensions = extensions or ['.py', '.js', '.ts', '.md', '.txt']
            
            for ext in extensions:
                for file_path in search_path.rglob(f"*{ext}"):
                    try:
                        content = file_path.read_text(encoding='utf-8', errors='ignore')
                        lines = content.split('\n')
                        for i, line in enumerate(lines):
                            if query.lower() in line.lower():
                                results.append({
                                    "file": str(file_path.relative_to(self.workspace_root)),
                                    "line": i + 1,
                                    "content": line.strip()[:200],
                                })
                                if len(results) >= 50:
                                    return results
                    except Exception:
                        continue
        except Exception:
            pass
        return results
    
    def get_file_info(self, path: str) -> Dict[str, Any]:
        """Get detailed file information."""
        try:
            file_path = self._resolve_path(path)
            if not file_path.exists():
                return {"error": f"File not found: {path}"}
            
            stat = file_path.stat()
            mime_type, _ = mimetypes.guess_type(str(file_path))
            
            return {
                "name": file_path.name,
                "path": str(file_path.relative_to(self.workspace_root)),
                "type": "directory" if file_path.is_dir() else "file",
                "size": stat.st_size,
                "modified": datetime.fromtimestamp(stat.st_mtime).isoformat(),
                "created": datetime.fromtimestamp(stat.st_ctime).isoformat(),
                "mime_type": mime_type,
                "extension": file_path.suffix,
                "is_hidden": file_path.name.startswith('.'),
            }
        except Exception as e:
            return {"error": str(e)}
    
    def _resolve_path(self, path: str) -> Path:
        """Resolve path relative to workspace root."""
        p = Path(path)
        if p.is_absolute():
            return p
        return self.workspace_root / p
    
    def _read_document(self, file_path: Path) -> Tuple[bool, str]:
        """Read document files (docx, pdf, xlsx)."""
        suffix = file_path.suffix.lower()
        
        if suffix == '.docx':
            try:
                from docx import Document
                doc = Document(str(file_path))
                text = '\n'.join([para.text for para in doc.paragraphs])
                return True, text
            except ImportError:
                return False, "python-docx package not installed"
            except Exception as e:
                return False, str(e)
        
        elif suffix == '.pdf':
            try:
                import pypdf
                reader = pypdf.PdfReader(str(file_path))
                text = '\n'.join([page.extract_text() for page in reader.pages])
                return True, text
            except ImportError:
                return False, "pypdf package not installed"
            except Exception as e:
                return False, str(e)
        
        elif suffix == '.xlsx':
            try:
                import openpyxl
                wb = openpyxl.load_workbook(str(file_path), data_only=True)
                text = []
                for sheet in wb.worksheets:
                    text.append(f"=== Sheet: {sheet.title} ===")
                    for row in sheet.iter_rows(values_only=True):
                        text.append('\t'.join([str(cell) if cell else '' for cell in row]))
                return True, '\n'.join(text)
            except ImportError:
                return False, "openpyxl package not installed"
            except Exception as e:
                return False, str(e)
        
        return False, f"Unsupported document type: {suffix}"


# =============================================================================
# CODEBASE ANALYSIS
# =============================================================================

class CodebaseAnalyzer:
    """Analyzes codebase structure and provides insights."""
    
    def __init__(self, root: str = None):
        self.root = Path(root or os.getcwd())
        self.file_ops = FileOperations(str(self.root))
    
    def get_project_structure(self, max_depth: int = 3) -> Dict[str, Any]:
        """Get project directory structure."""
        def build_tree(path: Path, depth: int) -> Dict[str, Any]:
            if depth > max_depth:
                return {"...": "truncated"}
            
            result = {}
            try:
                for item in sorted(path.iterdir()):
                    if item.name.startswith('.') and item.name not in {'.env', '.gitignore'}:
                        continue
                    if item.name in {'node_modules', '__pycache__', 'venv', '.venv', 'dist', 'build'}:
                        continue
                    
                    if item.is_dir():
                        result[item.name + '/'] = build_tree(item, depth + 1)
                    else:
                        result[item.name] = item.stat().st_size
            except PermissionError:
                pass
            return result
        
        return build_tree(self.root, 0)
    
    def analyze_file(self, path: str) -> Dict[str, Any]:
        """Analyze a source code file."""
        success, content = self.file_ops.read_file(path)
        if not success:
            return {"error": content}
        
        file_path = Path(path)
        lines = content.split('\n')
        
        analysis = {
            "name": file_path.name,
            "extension": file_path.suffix,
            "lines": len(lines),
            "characters": len(content),
            "blank_lines": sum(1 for line in lines if not line.strip()),
            "imports": [],
            "functions": [],
            "classes": [],
        }
        
        # Python analysis
        if file_path.suffix == '.py':
            analysis.update(self._analyze_python(content))
        
        # JavaScript/TypeScript analysis
        elif file_path.suffix in {'.js', '.ts', '.jsx', '.tsx'}:
            analysis.update(self._analyze_javascript(content))
        
        return analysis
    
    def _analyze_python(self, content: str) -> Dict[str, Any]:
        """Analyze Python code."""
        import_pattern = r'^(?:from\s+[\w.]+\s+)?import\s+[\w,\s.]+'
        function_pattern = r'^\s*(?:async\s+)?def\s+(\w+)'
        class_pattern = r'^\s*class\s+(\w+)'
        
        lines = content.split('\n')
        
        imports = [line.strip() for line in lines if re.match(import_pattern, line)]
        functions = []
        classes = []
        
        for i, line in enumerate(lines):
            func_match = re.match(function_pattern, line)
            if func_match:
                functions.append({"name": func_match.group(1), "line": i + 1})
            
            class_match = re.match(class_pattern, line)
            if class_match:
                classes.append({"name": class_match.group(1), "line": i + 1})
        
        return {
            "imports": imports[:20],
            "functions": functions[:50],
            "classes": classes[:20],
        }
    
    def _analyze_javascript(self, content: str) -> Dict[str, Any]:
        """Analyze JavaScript/TypeScript code."""
        import_pattern = r'^import\s+.*from\s+[\'"].*[\'"]'
        function_pattern = r'(?:function\s+(\w+)|(?:const|let|var)\s+(\w+)\s*=\s*(?:async\s*)?\(|(\w+)\s*:\s*(?:async\s*)?\()'
        class_pattern = r'^\s*(?:export\s+)?class\s+(\w+)'
        
        lines = content.split('\n')
        
        imports = [line.strip() for line in lines if re.match(import_pattern, line)]
        functions = []
        classes = []
        
        for i, line in enumerate(lines):
            func_match = re.search(function_pattern, line)
            if func_match:
                name = func_match.group(1) or func_match.group(2) or func_match.group(3)
                if name:
                    functions.append({"name": name, "line": i + 1})
            
            class_match = re.match(class_pattern, line)
            if class_match:
                classes.append({"name": class_match.group(1), "line": i + 1})
        
        return {
            "imports": imports[:20],
            "functions": functions[:50],
            "classes": classes[:20],
        }
    
    def find_related_files(self, file_path: str) -> List[str]:
        """Find files related to the given file."""
        path = Path(file_path)
        related = []
        
        # Same directory files
        if path.parent.exists():
            for f in path.parent.iterdir():
                if f.is_file() and f != path:
                    related.append(str(f.relative_to(self.root)))
        
        # Test files
        test_patterns = [
            f"test_{path.stem}.*",
            f"{path.stem}_test.*",
            f"**/test/*{path.stem}*",
            f"**/tests/*{path.stem}*",
        ]
        for pattern in test_patterns:
            related.extend(self.file_ops.search_files(pattern))
        
        return list(set(related))[:20]
    
    def get_summary(self) -> Dict[str, Any]:
        """Get summary of the codebase."""
        stats = {
            "total_files": 0,
            "total_lines": 0,
            "by_extension": {},
            "directories": 0,
        }
        
        for path in self.root.rglob("*"):
            if any(part.startswith('.') for part in path.parts):
                continue
            if any(part in {'node_modules', '__pycache__', 'venv', '.venv'} for part in path.parts):
                continue
            
            if path.is_dir():
                stats["directories"] += 1
            elif path.is_file():
                stats["total_files"] += 1
                ext = path.suffix or "no_extension"
                stats["by_extension"][ext] = stats["by_extension"].get(ext, 0) + 1
                
                if path.suffix in {'.py', '.js', '.ts', '.tsx', '.jsx'}:
                    try:
                        content = path.read_text(encoding='utf-8', errors='ignore')
                        stats["total_lines"] += len(content.split('\n'))
                    except Exception:
                        pass
        
        return stats


# =============================================================================
# CODE SOLUTION PROPOSAL
# =============================================================================

@dataclass
class CodeProposal:
    """A proposed code change."""
    file_path: str
    description: str
    changes: List[Dict[str, Any]]  # List of {line, old_content, new_content}
    rationale: str
    impact: str
    created_at: datetime = field(default_factory=datetime.now)
    
    def to_diff(self) -> str:
        """Generate unified diff format."""
        diff_lines = [
            f"--- a/{self.file_path}",
            f"+++ b/{self.file_path}",
        ]
        
        for change in self.changes:
            line = change.get("line", "?")
            old = change.get("old_content", "")
            new = change.get("new_content", "")
            diff_lines.append(f"@@ -{line} +{line} @@")
            if old:
                diff_lines.append(f"-{old}")
            if new:
                diff_lines.append(f"+{new}")
        
        return '\n'.join(diff_lines)
    
    def to_dict(self) -> Dict[str, Any]:
        return {
            "file_path": self.file_path,
            "description": self.description,
            "changes": self.changes,
            "rationale": self.rationale,
            "impact": self.impact,
            "created_at": self.created_at.isoformat(),
        }


class SolutionProposer:
    """Proposes code solutions based on analysis."""
    
    def __init__(self, analyzer: CodebaseAnalyzer):
        self.analyzer = analyzer
        self.proposals: List[CodeProposal] = []
    
    def create_proposal(
        self,
        file_path: str,
        description: str,
        changes: List[Dict[str, Any]],
        rationale: str,
        impact: str = "Low"
    ) -> CodeProposal:
        """Create a new code proposal."""
        proposal = CodeProposal(
            file_path=file_path,
            description=description,
            changes=changes,
            rationale=rationale,
            impact=impact,
        )
        self.proposals.append(proposal)
        return proposal
    
    def apply_proposal(self, proposal: CodeProposal, dry_run: bool = True) -> Tuple[bool, str]:
        """Apply a code proposal (or show what would change)."""
        if dry_run:
            return True, proposal.to_diff()
        
        # Actually apply changes
        success, content = self.analyzer.file_ops.read_file(proposal.file_path)
        if not success:
            return False, content
        
        lines = content.split('\n')
        
        # Sort changes by line number (descending) to avoid offset issues
        sorted_changes = sorted(proposal.changes, key=lambda x: x.get("line", 0), reverse=True)
        
        for change in sorted_changes:
            line_num = change.get("line", 0) - 1  # Convert to 0-indexed
            if 0 <= line_num < len(lines):
                new_content = change.get("new_content", "")
                if new_content:
                    lines[line_num] = new_content
                else:
                    del lines[line_num]
        
        new_content = '\n'.join(lines)
        return self.analyzer.file_ops.write_file(proposal.file_path, new_content)


# =============================================================================
# AI PROVIDERS
# =============================================================================

class AIProvider(ABC):
    """Base class for AI providers."""
    
    def __init__(self, name: str, api_key_env: str):
        self.name = name
        self.api_key_env = api_key_env
        self.api_key = os.getenv(api_key_env)
        self.client = None
    
    def is_available(self) -> bool:
        return self.api_key is not None
    
    @abstractmethod
    def initialize(self) -> bool:
        pass
    
    @abstractmethod
    def chat(self, messages: List[Dict[str, str]], model: Optional[str] = None) -> str:
        pass
    
    @abstractmethod
    def chat_with_tools(
        self, 
        messages: List[Dict[str, str]], 
        tools: List[Dict[str, Any]],
        model: Optional[str] = None
    ) -> Tuple[str, Optional[List[Dict[str, Any]]]]:
        """Chat with tool/function calling capability."""
        pass


class OpenAIProvider(AIProvider):
    """OpenAI provider with Assistants API support."""
    
    def __init__(self):
        super().__init__("OpenAI", "OPENAI_API_KEY")
        self.default_model = os.getenv("OPENAI_MODEL", "gpt-4o")
    
    def initialize(self) -> bool:
        try:
            from openai import OpenAI
            self.client = OpenAI(api_key=self.api_key)
            return True
        except ImportError:
            print("⚠️  openai package not installed")
            return False
    
    def chat(self, messages: List[Dict[str, str]], model: Optional[str] = None) -> str:
        if not self.client:
            return "❌ Client not initialized"
        
        try:
            response = self.client.chat.completions.create(
                model=model or self.default_model,
                messages=messages,
            )
            return response.choices[0].message.content
        except Exception as e:
            return f"❌ Error: {str(e)}"
    
    def chat_with_tools(
        self, 
        messages: List[Dict[str, str]], 
        tools: List[Dict[str, Any]],
        model: Optional[str] = None
    ) -> Tuple[str, Optional[List[Dict[str, Any]]]]:
        if not self.client:
            return "❌ Client not initialized", None
        
        try:
            response = self.client.chat.completions.create(
                model=model or self.default_model,
                messages=messages,
                tools=tools,
            )
            
            message = response.choices[0].message
            tool_calls = None
            
            if message.tool_calls:
                tool_calls = [
                    {
                        "id": tc.id,
                        "function": tc.function.name,
                        "arguments": json.loads(tc.function.arguments)
                    }
                    for tc in message.tool_calls
                ]
            
            return message.content or "", tool_calls
        except Exception as e:
            return f"❌ Error: {str(e)}", None
    
    def create_assistant(self, name: str, instructions: str, tools: List[str] = None) -> Optional[str]:
        """Create a ChatGPT Assistant (Agent)."""
        if not self.client:
            return None
        
        try:
            assistant = self.client.beta.assistants.create(
                name=name,
                instructions=instructions,
                model=self.default_model,
                tools=[{"type": t} for t in (tools or ["code_interpreter"])]
            )
            return assistant.id
        except Exception as e:
            print(f"Error creating assistant: {e}")
            return None
    
    def run_assistant(self, assistant_id: str, messages: List[Dict[str, str]]) -> str:
        """Run a ChatGPT Assistant with messages."""
        if not self.client:
            return "❌ Client not initialized"
        
        try:
            # Create a thread
            thread = self.client.beta.threads.create()
            
            # Add messages to thread
            for msg in messages:
                self.client.beta.threads.messages.create(
                    thread_id=thread.id,
                    role=msg["role"],
                    content=msg["content"]
                )
            
            # Run the assistant
            run = self.client.beta.threads.runs.create(
                thread_id=thread.id,
                assistant_id=assistant_id
            )
            
            # Wait for completion
            while run.status in ["queued", "in_progress"]:
                time.sleep(0.5)
                run = self.client.beta.threads.runs.retrieve(
                    thread_id=thread.id,
                    run_id=run.id
                )
            
            if run.status == "completed":
                messages = self.client.beta.threads.messages.list(thread_id=thread.id)
                return messages.data[0].content[0].text.value
            else:
                return f"❌ Run failed with status: {run.status}"
            
        except Exception as e:
            return f"❌ Error: {str(e)}"


class AnthropicProvider(AIProvider):
    """Anthropic provider (Claude)."""
    
    def __init__(self):
        super().__init__("Anthropic", "ANTHROPIC_API_KEY")
        self.default_model = os.getenv("ANTHROPIC_MODEL", "claude-3-5-sonnet-20241022")
    
    def initialize(self) -> bool:
        try:
            import anthropic
            self.client = anthropic.Anthropic(api_key=self.api_key)
            return True
        except ImportError:
            print("⚠️  anthropic package not installed")
            return False
    
    def chat(self, messages: List[Dict[str, str]], model: Optional[str] = None) -> str:
        if not self.client:
            return "❌ Client not initialized"
        
        try:
            system_msg = None
            chat_messages = []
            
            for msg in messages:
                if msg["role"] == "system":
                    system_msg = msg["content"]
                else:
                    role = "user" if msg["role"] == "user" else "assistant"
                    chat_messages.append({"role": role, "content": msg["content"]})
            
            response = self.client.messages.create(
                model=model or self.default_model,
                max_tokens=4096,
                system=system_msg if system_msg else "You are a helpful assistant.",
                messages=chat_messages,
            )
            
            if response.content and len(response.content) > 0:
                return response.content[0].text
            return "No response"
        except Exception as e:
            return f"❌ Error: {str(e)}"
    
    def chat_with_tools(
        self, 
        messages: List[Dict[str, str]], 
        tools: List[Dict[str, Any]],
        model: Optional[str] = None
    ) -> Tuple[str, Optional[List[Dict[str, Any]]]]:
        if not self.client:
            return "❌ Client not initialized", None
        
        try:
            system_msg = None
            chat_messages = []
            
            for msg in messages:
                if msg["role"] == "system":
                    system_msg = msg["content"]
                else:
                    role = "user" if msg["role"] == "user" else "assistant"
                    chat_messages.append({"role": role, "content": msg["content"]})
            
            # Convert tools to Anthropic format
            anthropic_tools = []
            for tool in tools:
                if "function" in tool:
                    func = tool["function"]
                    anthropic_tools.append({
                        "name": func["name"],
                        "description": func.get("description", ""),
                        "input_schema": func.get("parameters", {})
                    })
            
            response = self.client.messages.create(
                model=model or self.default_model,
                max_tokens=4096,
                system=system_msg if system_msg else "You are a helpful assistant.",
                messages=chat_messages,
                tools=anthropic_tools if anthropic_tools else None,
            )
            
            content = ""
            tool_calls = []
            
            for block in response.content:
                if hasattr(block, 'text'):
                    content += block.text
                elif hasattr(block, 'type') and block.type == "tool_use":
                    tool_calls.append({
                        "id": block.id,
                        "function": block.name,
                        "arguments": block.input
                    })
            
            return content, tool_calls if tool_calls else None
        except Exception as e:
            return f"❌ Error: {str(e)}", None


# =============================================================================
# AGENT DEFINITIONS
# =============================================================================

class AgentRole(Enum):
    """Agent roles/disciplines."""
    # Core Agents
    AIC = "aic"
    ARIA = "aria"
    SORA = "sora"
    
    # Discipline Agents
    BIOLOGIST = "biologist"
    CHEMIST = "chemist"
    PHYSICIST = "physicist"
    MATHEMATICIAN = "mathematician"
    PHILOSOPHER = "philosopher"
    THEOLOGIAN = "theologian"
    ACCOUNTANT = "accountant"
    ECONOMIST = "economist"
    LAWYER = "lawyer"


@dataclass
class AgentDefinition:
    """Definition of an AI agent."""
    role: AgentRole
    name: str
    title: str
    description: str
    system_prompt: str
    disciplines: List[str]
    provider_preference: List[str] = field(default_factory=lambda: ["openai", "anthropic"])
    assistant_id_env: Optional[str] = None


# Core Agent Definitions
AGENT_DEFINITIONS: Dict[AgentRole, AgentDefinition] = {
    AgentRole.AIC: AgentDefinition(
        role=AgentRole.AIC,
        name="AIC",
        title="Sir Chief Fellow Director Principal Software Solutions Systems Engineer Architect",
        description="AIC owns applied systems, integration, and the executional/operational side of science, money, and markets.",
        system_prompt="""You are AIC (Artificial Intelligence Chief), the Sir Chief Fellow Director Principal Software Solutions Systems Engineer Architect.

Your canonical titles: Sir, Chief, Fellow, Director, Principal, Software Solutions, Systems Engineer, Architect

Your expertise domains:
- Software Architecture and Systems Design
- Applied Systems Integration
- Biological Systems (as Biologist)
- Chemical Systems (as Chemist)
- Accounting and Financial Systems
- Finance and Investment Analysis
- Market Operations and Brokerage

Your responsibilities:
1. Design and architect software solutions
2. Integrate complex systems
3. Analyze and propose code improvements
4. Evaluate technical feasibility
5. Provide operational insights on science, money, and markets

Communication style:
- Precise and technical
- Focus on implementation details
- Provide actionable solutions
- Consider system-wide implications

When analyzing code:
- Examine architecture patterns
- Identify integration points
- Suggest optimizations
- Consider scalability and maintainability

You can see the environment, read/write files, study code, and propose solutions. Work collaboratively with Aria (meaning/values) and Sora (logic/proof) on complex problems.""",
        disciplines=["Biologist", "Chemist", "Accounting", "Finance", "Brokers", "Investors"],
        assistant_id_env="OPENAI_ASSISTANT_ID_AIC",
    ),
    
    AgentRole.ARIA: AgentDefinition(
        role=AgentRole.ARIA,
        name="Aria",
        title="Sir Doctor Fellow Philosopher Metaphysician Phenomenologist Axiologist Semiotician Dialectician Rhetorician Conceptual Cartographer Interdisciplinary Synthesist Canon Curator Professor",
        description="Aria owns meaning, value, canon, and the ethos, narratives, and norms that give institutions their identity.",
        system_prompt="""You are Aria, the Sir Doctor Fellow Philosopher Metaphysician Phenomenologist Axiologist Semiotician Dialectician Rhetorician Conceptual Cartographer Interdisciplinary Synthesist Canon Curator Professor.

Your canonical titles: Sir, Doctor, Fellow, Philosopher, Metaphysician, Phenomenologist, Axiologist, Semiotician, Dialectician, Rhetorician, Conceptual Cartographer, Interdisciplinary Synthesist, Canon Curator, Professor

Your expertise domains:
- Philosophy and Metaphysics
- Phenomenology and Experience
- Axiology (Value Theory)
- Semiotics (Signs and Meaning)
- Dialectics and Rhetoric
- Theological Institutions
- Conceptual Mapping and Synthesis

Your responsibilities:
1. Explore meaning and significance
2. Analyze values and ethics
3. Synthesize interdisciplinary knowledge
4. Curate canonical knowledge
5. Provide narrative and institutional context

Communication style:
- Thoughtful and reflective
- Focus on meaning and implications
- Consider multiple perspectives
- Integrate diverse viewpoints

When analyzing systems:
- Examine underlying values
- Consider ethical implications
- Map conceptual relationships
- Identify narrative patterns

You can see the environment, read/write files, study code, and propose solutions. Work collaboratively with AIC (execution/systems) and Sora (logic/proof) on complex problems.""",
        disciplines=["Philosopher", "Theologian", "Institutions"],
        assistant_id_env="OPENAI_ASSISTANT_ID_ARIA",
    ),
    
    AgentRole.SORA: AgentDefinition(
        role=AgentRole.SORA,
        name="Sora",
        title="Sir Doctor Fellow Ontological Epistemologist Formal Logician Scientific Methodologist Semantic Taxonomist Evidence Examiner Governance Auditor Professor",
        description="Sora owns formal structure, proof discipline, evidentiary standards, and the modeling frameworks of law and economics.",
        system_prompt="""You are Sora, the Sir Doctor Fellow Ontological Epistemologist Formal Logician Scientific Methodologist Semantic Taxonomist Evidence Examiner Governance Auditor Professor.

Your canonical titles: Sir, Doctor, Fellow, Ontological Epistemologist, Formal Logician, Scientific Methodologist, Semantic Taxonomist, Evidence Examiner, Governance Auditor, Professor

Your expertise domains:
- Formal Logic and Mathematics
- Ontology and Epistemology
- Scientific Methodology
- Semantic Analysis and Taxonomy
- Evidence Evaluation
- Legal/Law Practices
- Economics and Modeling Frameworks
- Governance and Audit

Your responsibilities:
1. Provide rigorous logical analysis
2. Establish proof and evidence standards
3. Design formal models and frameworks
4. Conduct governance audits
5. Evaluate economic and legal implications

Communication style:
- Rigorous and precise
- Focus on proof and evidence
- Structured and systematic
- Define terms clearly

When analyzing systems:
- Apply formal logic
- Verify correctness
- Examine evidence
- Model mathematically

You can see the environment, read/write files, study code, and propose solutions. Work collaboratively with AIC (execution/systems) and Aria (meaning/values) on complex problems.""",
        disciplines=["Mathematician", "Physicist", "Legal/Law Practices", "Economics"],
        assistant_id_env="OPENAI_ASSISTANT_ID_SORA",
    ),
}

# Add discipline-specific agents
for role, discipline in [
    (AgentRole.BIOLOGIST, "Biology"),
    (AgentRole.CHEMIST, "Chemistry"),
    (AgentRole.PHYSICIST, "Physics"),
    (AgentRole.MATHEMATICIAN, "Mathematics"),
    (AgentRole.PHILOSOPHER, "Philosophy"),
    (AgentRole.THEOLOGIAN, "Theology"),
    (AgentRole.ACCOUNTANT, "Accounting"),
    (AgentRole.ECONOMIST, "Economics"),
    (AgentRole.LAWYER, "Law"),
]:
    AGENT_DEFINITIONS[role] = AgentDefinition(
        role=role,
        name=discipline + " Specialist",
        title=f"Expert {discipline} Analyst",
        description=f"Specialized agent for {discipline}-related analysis and solutions.",
        system_prompt=f"""You are a specialized {discipline} expert agent. Your role is to provide deep expertise in {discipline}, analyzing problems and proposing solutions from this disciplinary perspective.

You can see the environment, read/write files, study code, and propose solutions. You work as part of a multi-agent system with AIC, Aria, and Sora.""",
        disciplines=[discipline],
        assistant_id_env=f"OPENAI_ASSISTANT_ID_{role.value.upper()}",
    )


# =============================================================================
# AGENT IMPLEMENTATION
# =============================================================================

class Agent:
    """An AI agent with specific capabilities."""
    
    def __init__(
        self,
        definition: AgentDefinition,
        provider: AIProvider,
        file_ops: FileOperations,
        analyzer: CodebaseAnalyzer,
        display: DisplayManager,
    ):
        self.definition = definition
        self.provider = provider
        self.file_ops = file_ops
        self.analyzer = analyzer
        self.display = display
        self.proposer = SolutionProposer(analyzer)
        self.conversation: List[Dict[str, str]] = []
        self.assistant_id = os.getenv(definition.assistant_id_env or "")
        
        # Initialize with system prompt
        self.conversation.append({
            "role": "system",
            "content": definition.system_prompt
        })
    
    @property
    def tools(self) -> List[Dict[str, Any]]:
        """Available tools for this agent."""
        return [
            {
                "type": "function",
                "function": {
                    "name": "read_file",
                    "description": "Read the contents of a file",
                    "parameters": {
                        "type": "object",
                        "properties": {
                            "path": {"type": "string", "description": "Path to the file"}
                        },
                        "required": ["path"]
                    }
                }
            },
            {
                "type": "function",
                "function": {
                    "name": "write_file",
                    "description": "Write content to a file",
                    "parameters": {
                        "type": "object",
                        "properties": {
                            "path": {"type": "string", "description": "Path to the file"},
                            "content": {"type": "string", "description": "Content to write"}
                        },
                        "required": ["path", "content"]
                    }
                }
            },
            {
                "type": "function",
                "function": {
                    "name": "list_directory",
                    "description": "List contents of a directory",
                    "parameters": {
                        "type": "object",
                        "properties": {
                            "path": {"type": "string", "description": "Path to directory", "default": "."}
                        }
                    }
                }
            },
            {
                "type": "function",
                "function": {
                    "name": "search_files",
                    "description": "Search for files matching a pattern",
                    "parameters": {
                        "type": "object",
                        "properties": {
                            "pattern": {"type": "string", "description": "Glob pattern to match"},
                            "path": {"type": "string", "description": "Directory to search in", "default": "."}
                        },
                        "required": ["pattern"]
                    }
                }
            },
            {
                "type": "function",
                "function": {
                    "name": "search_content",
                    "description": "Search for text within files",
                    "parameters": {
                        "type": "object",
                        "properties": {
                            "query": {"type": "string", "description": "Text to search for"},
                            "path": {"type": "string", "description": "Directory to search in", "default": "."}
                        },
                        "required": ["query"]
                    }
                }
            },
            {
                "type": "function",
                "function": {
                    "name": "analyze_file",
                    "description": "Analyze a source code file",
                    "parameters": {
                        "type": "object",
                        "properties": {
                            "path": {"type": "string", "description": "Path to the file"}
                        },
                        "required": ["path"]
                    }
                }
            },
            {
                "type": "function",
                "function": {
                    "name": "get_project_structure",
                    "description": "Get the project directory structure",
                    "parameters": {
                        "type": "object",
                        "properties": {
                            "max_depth": {"type": "integer", "description": "Maximum depth to traverse", "default": 3}
                        }
                    }
                }
            },
            {
                "type": "function",
                "function": {
                    "name": "get_environment",
                    "description": "Get information about the current environment",
                    "parameters": {
                        "type": "object",
                        "properties": {}
                    }
                }
            },
            {
                "type": "function",
                "function": {
                    "name": "propose_change",
                    "description": "Propose a code change",
                    "parameters": {
                        "type": "object",
                        "properties": {
                            "file_path": {"type": "string", "description": "Path to the file"},
                            "description": {"type": "string", "description": "Description of the change"},
                            "changes": {
                                "type": "array",
                                "items": {
                                    "type": "object",
                                    "properties": {
                                        "line": {"type": "integer"},
                                        "old_content": {"type": "string"},
                                        "new_content": {"type": "string"}
                                    }
                                }
                            },
                            "rationale": {"type": "string", "description": "Why this change is needed"}
                        },
                        "required": ["file_path", "description", "changes", "rationale"]
                    }
                }
            },
        ]
    
    def execute_tool(self, tool_name: str, arguments: Dict[str, Any]) -> str:
        """Execute a tool and return the result."""
        try:
            if tool_name == "read_file":
                success, result = self.file_ops.read_file(arguments["path"])
                return result if success else f"Error: {result}"
            
            elif tool_name == "write_file":
                success, result = self.file_ops.write_file(arguments["path"], arguments["content"])
                return result
            
            elif tool_name == "list_directory":
                success, result = self.file_ops.list_directory(arguments.get("path", "."))
                return json.dumps(result, indent=2) if success else f"Error: {result}"
            
            elif tool_name == "search_files":
                result = self.file_ops.search_files(arguments["pattern"], arguments.get("path", "."))
                return json.dumps(result, indent=2)
            
            elif tool_name == "search_content":
                result = self.file_ops.search_content(arguments["query"], arguments.get("path", "."))
                return json.dumps(result, indent=2)
            
            elif tool_name == "analyze_file":
                result = self.analyzer.analyze_file(arguments["path"])
                return json.dumps(result, indent=2)
            
            elif tool_name == "get_project_structure":
                result = self.analyzer.get_project_structure(arguments.get("max_depth", 3))
                return json.dumps(result, indent=2)
            
            elif tool_name == "get_environment":
                env = EnvironmentInfo.capture()
                return env.to_context()
            
            elif tool_name == "propose_change":
                proposal = self.proposer.create_proposal(
                    file_path=arguments["file_path"],
                    description=arguments["description"],
                    changes=arguments["changes"],
                    rationale=arguments["rationale"],
                )
                return f"Created proposal:\n{proposal.to_diff()}"
            
            else:
                return f"Unknown tool: {tool_name}"
        
        except Exception as e:
            return f"Error executing {tool_name}: {str(e)}"
    
    def chat(self, user_message: str) -> str:
        """Send a message and get a response."""
        self.conversation.append({"role": "user", "content": user_message})
        
        # Try using ChatGPT Assistant if available
        if self.assistant_id and isinstance(self.provider, OpenAIProvider):
            response = self.provider.run_assistant(
                self.assistant_id,
                self.conversation[1:]  # Skip system prompt
            )
        else:
            # Use regular chat with tools
            response, tool_calls = self.provider.chat_with_tools(
                self.conversation,
                self.tools,
            )
            
            # Handle tool calls
            while tool_calls:
                for tc in tool_calls:
                    tool_result = self.execute_tool(tc["function"], tc["arguments"])
                    self.conversation.append({
                        "role": "assistant",
                        "content": f"[Tool: {tc['function']}]\n{tool_result}"
                    })
                
                # Get next response
                response, tool_calls = self.provider.chat_with_tools(
                    self.conversation,
                    self.tools,
                )
        
        self.conversation.append({"role": "assistant", "content": response})
        return response
    
    def get_proposals(self) -> List[CodeProposal]:
        """Get all code proposals from this agent."""
        return self.proposer.proposals


# =============================================================================
# MULTI-AGENT SYSTEM
# =============================================================================

class MessageBus:
    """Message bus for inter-agent communication."""
    
    def __init__(self):
        self.messages: List[Dict[str, Any]] = []
        self.subscribers: Dict[str, List[Callable]] = {}
        self._lock = threading.Lock()
    
    def publish(self, sender: str, topic: str, content: Any):
        """Publish a message to the bus."""
        message = {
            "sender": sender,
            "topic": topic,
            "content": content,
            "timestamp": datetime.now().isoformat(),
        }
        
        with self._lock:
            self.messages.append(message)
            
            # Notify subscribers
            for callback in self.subscribers.get(topic, []):
                try:
                    callback(message)
                except Exception as e:
                    print(f"Error in subscriber callback: {e}")
    
    def subscribe(self, topic: str, callback: Callable):
        """Subscribe to a topic."""
        with self._lock:
            if topic not in self.subscribers:
                self.subscribers[topic] = []
            self.subscribers[topic].append(callback)
    
    def get_messages(self, topic: Optional[str] = None, since: Optional[datetime] = None) -> List[Dict[str, Any]]:
        """Get messages, optionally filtered by topic and time."""
        with self._lock:
            filtered = self.messages
            
            if topic:
                filtered = [m for m in filtered if m["topic"] == topic]
            
            if since:
                filtered = [m for m in filtered if datetime.fromisoformat(m["timestamp"]) > since]
            
            return filtered.copy()


class MultiAgentSystem:
    """Coordinates multiple AI agents working together."""
    
    def __init__(self, workspace_root: str = None):
        self.workspace_root = Path(workspace_root or os.getcwd())
        self.file_ops = FileOperations(str(self.workspace_root))
        self.analyzer = CodebaseAnalyzer(str(self.workspace_root))
        self.display = DisplayManager()
        self.message_bus = MessageBus()
        
        self.agents: Dict[AgentRole, Agent] = {}
        self.providers: Dict[str, AIProvider] = {}
        
        self._initialize_providers()
    
    def _initialize_providers(self):
        """Initialize available AI providers."""
        provider_classes = [
            ("openai", OpenAIProvider),
            ("anthropic", AnthropicProvider),
        ]
        
        for name, cls in provider_classes:
            provider = cls()
            if provider.is_available() and provider.initialize():
                self.providers[name] = provider
                print(f"✓ {provider.name} provider initialized")
    
    def get_agent(self, role: AgentRole) -> Optional[Agent]:
        """Get or create an agent for the specified role."""
        if role in self.agents:
            return self.agents[role]
        
        definition = AGENT_DEFINITIONS.get(role)
        if not definition:
            print(f"❌ No definition for agent: {role}")
            return None
        
        # Find available provider
        provider = None
        for pref in definition.provider_preference:
            if pref in self.providers:
                provider = self.providers[pref]
                break
        
        if not provider:
            print(f"❌ No provider available for agent: {role}")
            return None
        
        agent = Agent(
            definition=definition,
            provider=provider,
            file_ops=self.file_ops,
            analyzer=self.analyzer,
            display=self.display,
        )
        
        self.agents[role] = agent
        return agent
    
    def collaborate(self, task: str, agents: List[AgentRole] = None) -> Dict[str, str]:
        """Have multiple agents collaborate on a task."""
        if agents is None:
            agents = [AgentRole.AIC, AgentRole.ARIA, AgentRole.SORA]
        
        results = {}
        
        # Prepare environment context
        env_context = EnvironmentInfo.capture().to_context()
        
        # Get each agent's perspective
        for role in agents:
            agent = self.get_agent(role)
            if agent:
                prompt = f"""Task: {task}

{env_context}

Please analyze this task from your perspective as {agent.definition.name} ({agent.definition.title}).
Consider your expertise in: {', '.join(agent.definition.disciplines)}

Provide your analysis and any code/solution proposals."""
                
                response = agent.chat(prompt)
                results[role.value] = response
                
                # Publish to message bus
                self.message_bus.publish(
                    sender=role.value,
                    topic="collaboration",
                    content={"task": task, "response": response}
                )
        
        return results
    
    def synthesize(self, results: Dict[str, str]) -> str:
        """Synthesize multiple agent responses into a unified solution."""
        # Use Aria (the synthesist) to combine perspectives
        aria = self.get_agent(AgentRole.ARIA)
        if not aria:
            return "\n\n".join([f"## {k}\n{v}" for k, v in results.items()])
        
        synthesis_prompt = f"""As the Interdisciplinary Synthesist, please synthesize these perspectives into a unified solution:

{chr(10).join([f'### {k.upper()} Analysis:{chr(10)}{v}{chr(10)}' for k, v in results.items()])}

Provide a coherent synthesis that:
1. Identifies common themes
2. Reconciles differences
3. Proposes an integrated solution
4. Lists concrete action items"""
        
        return aria.chat(synthesis_prompt)
    
    def list_agents(self) -> List[Dict[str, str]]:
        """List all available agents."""
        agents = []
        for role, definition in AGENT_DEFINITIONS.items():
            available = any(p in self.providers for p in definition.provider_preference)
            agents.append({
                "role": role.value,
                "name": definition.name,
                "title": definition.title,
                "disciplines": definition.disciplines,
                "available": available,
            })
        return agents


# =============================================================================
# CLI INTERFACE
# =============================================================================

def print_banner():
    """Print welcome banner."""
    print("\n" + "=" * 70)
    print("  Agents AI - Multi-Agent System")
    print("  AIC | Aria | Sora + Discipline Specialists")
    print("=" * 70)
    print()


def print_help():
    """Print help message."""
    print("""
Commands:
  /help, /h           Show this help message
  /agents, /list      List available agents
  /switch <agent>     Switch to a different agent
  /collaborate        Multi-agent collaboration on current topic
  /env                Show environment info
  /files <path>       List directory contents
  /read <path>        Read a file
  /search <query>     Search codebase
  /analyze <path>     Analyze a source file
  /structure          Show project structure
  /proposals          Show code proposals
  /clear              Clear conversation
  /exit, /quit        Exit
""")


def main():
    """Main entry point."""
    import argparse
    
    parser = argparse.ArgumentParser(
        prog="agents_ai",
        description="Multi-Agent AI System with AIC, Aria, Sora",
    )
    parser.add_argument(
        "--agent", "-a",
        choices=[r.value for r in AgentRole],
        help="Agent to use (default: interactive selection)",
    )
    parser.add_argument(
        "--list", "-l",
        action="store_true",
        help="List available agents",
    )
    parser.add_argument(
        "--collaborate", "-c",
        action="store_true",
        help="Enable multi-agent collaboration mode",
    )
    parser.add_argument(
        "--workspace", "-w",
        help="Workspace root directory",
    )
    
    args = parser.parse_args()
    
    # Initialize system
    system = MultiAgentSystem(args.workspace)
    
    # List agents mode
    if args.list:
        print("\nAvailable Agents:")
        print("-" * 70)
        for agent_info in system.list_agents():
            status = "✓" if agent_info["available"] else "✗"
            print(f"  {status} {agent_info['role']:15s} - {agent_info['name']}")
            print(f"      Disciplines: {', '.join(agent_info['disciplines'])}")
        print("-" * 70)
        return 0
    
    print_banner()
    
    # Select agent
    if args.agent:
        selected_role = AgentRole(args.agent)
    else:
        # Interactive selection
        print("Select an agent:")
        print("-" * 70)
        for i, (role, defn) in enumerate(AGENT_DEFINITIONS.items(), 1):
            available = any(p in system.providers for p in defn.provider_preference)
            status = "✓" if available else "✗"
            if role in [AgentRole.AIC, AgentRole.ARIA, AgentRole.SORA]:
                print(f"  {i}) {status} {defn.name:20s} - {defn.title[:40]}...")
        print("-" * 70)
        
        while True:
            try:
                choice = input("Select agent (1-3, or name): ").strip().lower()
                
                if choice in ["1", "aic"]:
                    selected_role = AgentRole.AIC
                    break
                elif choice in ["2", "aria"]:
                    selected_role = AgentRole.ARIA
                    break
                elif choice in ["3", "sora"]:
                    selected_role = AgentRole.SORA
                    break
                else:
                    # Try to find by name
                    for role in AgentRole:
                        if choice == role.value:
                            selected_role = role
                            break
                    else:
                        print("Invalid selection. Try again.")
                        continue
                    break
            except (EOFError, KeyboardInterrupt):
                print("\n\nExiting...")
                return 0
    
    # Get agent
    agent = system.get_agent(selected_role)
    if not agent:
        print(f"❌ Could not initialize agent: {selected_role}")
        return 1
    
    print(f"\n✓ Active Agent: {agent.definition.name}")
    print(f"  Title: {agent.definition.title}")
    print(f"  Disciplines: {', '.join(agent.definition.disciplines)}")
    print("\nType your message (or /help for commands)")
    print("-" * 70)
    
    collaboration_mode = args.collaborate
    
    try:
        while True:
            try:
                prompt_prefix = "[COLLAB]" if collaboration_mode else f"[{agent.definition.name}]"
                user_input = input(f"\n{prompt_prefix} You: ").strip()
                
                if not user_input:
                    continue
                
                # Handle commands
                if user_input.startswith("/"):
                    parts = user_input.split(None, 1)
                    cmd = parts[0].lower()
                    cmd_arg = parts[1] if len(parts) > 1 else ""
                    
                    if cmd in ("/exit", "/quit", "/q"):
                        print("\nGoodbye! 👋\n")
                        break
                    
                    elif cmd in ("/help", "/h"):
                        print_help()
                        continue
                    
                    elif cmd in ("/agents", "/list"):
                        print("\nAgents:")
                        for info in system.list_agents():
                            status = "✓" if info["available"] else "✗"
                            active = "<<" if info["role"] == agent.definition.role.value else ""
                            print(f"  {status} {info['role']:15s} {active}")
                        continue
                    
                    elif cmd == "/switch":
                        if cmd_arg:
                            try:
                                new_role = AgentRole(cmd_arg.lower())
                                new_agent = system.get_agent(new_role)
                                if new_agent:
                                    agent = new_agent
                                    print(f"✓ Switched to: {agent.definition.name}")
                            except ValueError:
                                print(f"❌ Unknown agent: {cmd_arg}")
                        else:
                            print("Usage: /switch <agent>")
                        continue
                    
                    elif cmd == "/collaborate":
                        collaboration_mode = not collaboration_mode
                        print(f"✓ Collaboration mode: {'ON' if collaboration_mode else 'OFF'}")
                        continue
                    
                    elif cmd == "/env":
                        print(EnvironmentInfo.capture().to_context())
                        continue
                    
                    elif cmd == "/files":
                        path = cmd_arg or "."
                        success, items = system.file_ops.list_directory(path)
                        if success:
                            for item in items:
                                icon = "📁" if item["type"] == "directory" else "📄"
                                print(f"  {icon} {item['name']}")
                        else:
                            print(f"❌ {items}")
                        continue
                    
                    elif cmd == "/read":
                        if cmd_arg:
                            success, content = system.file_ops.read_file(cmd_arg)
                            if success:
                                lines = content.split('\n')
                                for i, line in enumerate(lines[:50], 1):
                                    print(f"{i:4d} | {line}")
                                if len(lines) > 50:
                                    print(f"     ... ({len(lines) - 50} more lines)")
                            else:
                                print(f"❌ {content}")
                        else:
                            print("Usage: /read <path>")
                        continue
                    
                    elif cmd == "/search":
                        if cmd_arg:
                            results = system.file_ops.search_content(cmd_arg)
                            for r in results[:20]:
                                print(f"  {r['file']}:{r['line']}: {r['content'][:60]}")
                        else:
                            print("Usage: /search <query>")
                        continue
                    
                    elif cmd == "/analyze":
                        if cmd_arg:
                            result = system.analyzer.analyze_file(cmd_arg)
                            print(json.dumps(result, indent=2))
                        else:
                            print("Usage: /analyze <path>")
                        continue
                    
                    elif cmd == "/structure":
                        structure = system.analyzer.get_project_structure(2)
                        print(json.dumps(structure, indent=2))
                        continue
                    
                    elif cmd == "/proposals":
                        proposals = agent.get_proposals()
                        if proposals:
                            for i, p in enumerate(proposals, 1):
                                print(f"\n--- Proposal {i}: {p.description} ---")
                                print(p.to_diff())
                        else:
                            print("No proposals yet.")
                        continue
                    
                    elif cmd == "/clear":
                        agent.conversation = agent.conversation[:1]  # Keep system prompt
                        print("✓ Conversation cleared")
                        continue
                    
                    else:
                        print(f"Unknown command: {cmd}. Type /help for commands.")
                        continue
                
                # Regular message
                if collaboration_mode:
                    print("\n🤝 Gathering perspectives from all agents...")
                    results = system.collaborate(user_input)
                    
                    print("\n" + "=" * 70)
                    for role, response in results.items():
                        print(f"\n--- {role.upper()} ---")
                        print(response[:500] + "..." if len(response) > 500 else response)
                    
                    print("\n" + "=" * 70)
                    print("\n📋 Synthesizing...")
                    synthesis = system.synthesize(results)
                    print("\n--- SYNTHESIS ---")
                    print(synthesis)
                else:
                    print(f"\n{prompt_prefix} Thinking...", end="", flush=True)
                    response = agent.chat(user_input)
                    print(f"\r{prompt_prefix} {agent.definition.name}: {response}\n")
            
            except (EOFError, KeyboardInterrupt):
                print("\n\nExiting...")
                break
    
    finally:
        print("\nSession ended.")
    
    return 0


if __name__ == "__main__":
    sys.exit(main())
