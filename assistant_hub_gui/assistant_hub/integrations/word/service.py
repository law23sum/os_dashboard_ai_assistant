"""Word service for drafting and rewriting documents with AI."""

from __future__ import annotations

from pathlib import Path
from typing import List

from docx import Document

from ...ai_layer.tools import summarize_text
from .local_client import load_document, save_document


class WordService:
    """High-level Word document operations."""

    def draft_summary(self, source_text: str, destination: str, actor: str = "Aria") -> List[str]:
        summary = summarize_text(source_text, style="clear")
        doc = Document()
        doc.add_paragraph(summary)
        dest_path = Path(destination)
        dest_path.parent.mkdir(parents=True, exist_ok=True)
        save_document(doc, destination)
        return [str(dest_path)]

    def rewrite_document(self, path: str, style: str = "concise", actor: str = "Aria") -> List[str]:
        doc = load_document(path)
        paragraphs = "\n".join(p.text for p in doc.paragraphs)
        rewritten = summarize_text(paragraphs, style=style)
        new_doc = Document()
        for line in rewritten.split("\n"):
            new_doc.add_paragraph(line)
        save_document(new_doc, path)
        return [path]
